#!/usr/bin/env python3
"""
Create final cleaned Word document - removes all prompt hallucinations
"""

import json
import re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def format_timestamp(seconds):
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"

def correct_names(text):
    corrections = {
        r'\bRaleigh\b': 'Rylie',
        r'\bRayleigh\b': 'Rylie', 
        r'\bRylee\b': 'Rylie',
        r'\bRiley\b': 'Rylie',
        r'\bMarcello\b': 'Marcella',
        r'\bMarsella\b': 'Marcella',
        r'\bScronse\b': 'Scronce',
        r'\bSkronce\b': 'Scronce',
        r'\bBlankship\b': 'Blankenship',
    }
    result = text
    for pattern, replacement in corrections.items():
        result = re.sub(pattern, replacement, result, flags=re.IGNORECASE)
    return result

def is_hallucination(text):
    """Detect prompt hallucinations and noise"""
    text_lower = text.strip().lower()
    
    # Remove R-Y-L-I-E patterns and check what's left
    cleaned = re.sub(r'r-y-l-i-e[,.\s]*', '', text_lower).strip()
    cleaned = re.sub(r'transcribe verbatim[,.\s]*', '', cleaned).strip()
    cleaned = re.sub(r'police interrogation[,.\s]*', '', cleaned).strip()
    
    # If nothing substantial left, it's noise
    if len(cleaned) < 5:
        return True
    
    # Check for mostly spelling repetitions
    rylie_count = text_lower.count('r-y-l-i-e')
    if rylie_count > 2:
        return True
    
    # Check for repeated "please" patterns
    if text_lower.count('mr. blankenship, please') > 1:
        return True
    
    return False

def main():
    with open('/app/backend/transcription/transcript_segments.json') as f:
        segments = json.load(f)
    
    print(f"Original segments: {len(segments)}")
    
    # First pass - identify silence regions
    silence_regions = []
    i = 0
    while i < len(segments):
        # Check for hallucination clusters
        cluster_start = i
        while i < len(segments) and is_hallucination(segments[i]['text']):
            i += 1
        
        if i - cluster_start >= 5:  # 5+ consecutive hallucinations = silence
            silence_regions.append({
                'start': segments[cluster_start]['start'],
                'end': segments[min(i-1, len(segments)-1)]['start']
            })
        i += 1
    
    print(f"Found {len(silence_regions)} silence regions")
    
    # Second pass - build clean segments
    clean_segments = []
    silence_idx = 0
    
    for seg in segments:
        # Check if in a silence region
        in_silence = False
        for sr in silence_regions:
            if sr['start'] <= seg['start'] <= sr['end']:
                in_silence = True
                break
        
        if in_silence:
            continue
        
        # Skip hallucinations
        if is_hallucination(seg['text']):
            continue
        
        # Skip very short segments
        if len(seg['text'].strip()) < 3:
            continue
        
        clean_segments.append(seg)
    
    # Add silence markers
    final_segments = []
    last_end = 0
    
    for sr in silence_regions:
        # Add segments before this silence
        for seg in clean_segments:
            if last_end <= seg['start'] < sr['start']:
                final_segments.append(seg)
        
        # Add silence marker if gap > 30 seconds
        if sr['end'] - sr['start'] > 30:
            final_segments.append({
                'start': sr['start'],
                'text': f"[Silence - approximately {int((sr['end']-sr['start'])/60)} minutes]",
                'is_silence': True
            })
        
        last_end = sr['end']
    
    # Add remaining segments after last silence
    for seg in clean_segments:
        if seg['start'] >= last_end:
            final_segments.append(seg)
    
    # Sort by timestamp
    final_segments.sort(key=lambda x: x['start'])
    
    print(f"Final segments: {len(final_segments)}")
    
    # Create Word document
    doc = Document()
    
    title = doc.add_heading('Catawba County Interrogation Transcript', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.add_run("Case: ").bold = True
    meta.add_run("State v. Zachary Blankenship\n")
    meta.add_run("Date: ").bold = True
    meta.add_run("December 13, 2013\n")
    meta.add_run("Time: ").bold = True
    meta.add_run("Approximately 11:18 AM\n")
    meta.add_run("Location: ").bold = True
    meta.add_run("Catawba County\n")
    meta.add_run("Total Recording Duration: ").bold = True
    meta.add_run("02:45:27\n")
    meta.add_run("\nParticipants:\n").bold = True
    meta.add_run("• Marcella McCombs - Lead Interrogator (Female)\n")
    meta.add_run("• Sergeant Scronce - Interrogator (Male)\n")
    meta.add_run("• Zachary Blankenship - Defendant\n")
    meta.add_run("• Rylie Blankenship - Victim (referenced)\n")
    
    doc.add_paragraph()
    doc.add_paragraph("─" * 60)
    doc.add_heading("VERBATIM TRANSCRIPT", level=1)
    doc.add_paragraph()
    
    # Track to avoid duplicates
    seen_texts = set()
    last_time = -5
    
    for seg in final_segments:
        timestamp = format_timestamp(seg['start'])
        text = correct_names(seg['text'])
        
        if not text.strip():
            continue
        
        # Skip near-duplicates
        text_key = text.strip().lower()[:50]
        if text_key in seen_texts and seg['start'] - last_time < 3:
            continue
        seen_texts.add(text_key)
        last_time = seg['start']
        
        p = doc.add_paragraph()
        
        ts_run = p.add_run(f"[{timestamp}] ")
        ts_run.bold = True
        ts_run.font.size = Pt(10)
        
        text_run = p.add_run(text)
        text_run.font.size = Pt(11)
    
    output_path = '/app/backend/transcription/Catawba_Interrogation_FINAL.docx'
    doc.save(output_path)
    print(f"\nSaved: {output_path}")

if __name__ == "__main__":
    main()
