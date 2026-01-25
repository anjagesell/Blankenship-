#!/usr/bin/env python3
"""
Transcription script for the Catawba interrogation audio.
Downloads, splits, transcribes with OpenAI Whisper, and creates Word doc.
"""

import os
import asyncio
import tempfile
import subprocess
from pathlib import Path
from dotenv import load_dotenv
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re
import json

load_dotenv("/app/backend/.env")

def format_timestamp(seconds):
    """Convert seconds to HH:MM:SS format"""
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"

def split_audio_ffmpeg(input_path, output_dir, chunk_duration=600):
    """Split audio into chunks using ffmpeg"""
    chunks = []
    
    result = subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration", 
         "-of", "default=noprint_wrappers=1:nokey=1", input_path],
        capture_output=True, text=True
    )
    total_duration = float(result.stdout.strip())
    print(f"Total duration: {format_timestamp(total_duration)}")
    
    chunk_num = 0
    start = 0
    
    while start < total_duration:
        chunk_path = os.path.join(output_dir, f"chunk_{chunk_num:03d}.mp3")
        
        subprocess.run([
            "ffmpeg", "-y", "-i", input_path,
            "-ss", str(start), "-t", str(chunk_duration),
            "-acodec", "libmp3lame", "-b:a", "64k",
            chunk_path
        ], capture_output=True)
        
        chunks.append({
            "path": chunk_path,
            "start_time": start,
            "chunk_num": chunk_num
        })
        print(f"Chunk {chunk_num}: {format_timestamp(start)}")
        
        start += chunk_duration
        chunk_num += 1
    
    return chunks, total_duration

async def transcribe_chunk(api_key, chunk_info, prompt):
    """Transcribe a single audio chunk"""
    from emergentintegrations.llm.openai import OpenAISpeechToText
    
    stt = OpenAISpeechToText(api_key=api_key)
    chunk_path = chunk_info["path"]
    start_offset = chunk_info["start_time"]
    
    print(f"  Transcribing chunk {chunk_info['chunk_num']}...")
    
    try:
        with open(chunk_path, "rb") as audio_file:
            response = await stt.transcribe(
                file=audio_file,
                model="whisper-1",
                response_format="verbose_json",
                language="en",
                prompt=prompt,
                temperature=0.0,
                timestamp_granularities=["segment"]
            )
        
        segments = []
        if hasattr(response, 'segments') and response.segments:
            for seg in response.segments:
                segments.append({
                    "start": seg.start + start_offset,
                    "end": seg.end + start_offset,
                    "text": seg.text.strip()
                })
            print(f"    Got {len(segments)} segments")
        elif response.text:
            segments.append({
                "start": start_offset,
                "end": start_offset + 600,
                "text": response.text.strip()
            })
            print(f"    Got full text response")
        
        return segments
        
    except Exception as e:
        print(f"  ERROR chunk {chunk_info['chunk_num']}: {e}")
        return []

def correct_names(text):
    """Correct transcription errors in names"""
    corrections = {
        r'\bRaleigh\b': 'Rylie',
        r'\bRayleigh\b': 'Rylie', 
        r'\bRylee\b': 'Rylie',
        r'\bRiley\b': 'Rylie',
        r'\bMarcello\b': 'Marcella',
        r'\bMarsella\b': 'Marcella',
        r'\bScronse\b': 'Scronce',
        r'\bSkronce\b': 'Scronce',
        r'\bBlankship\b': 'Blankenship',
    }
    
    result = text
    for pattern, replacement in corrections.items():
        result = re.sub(pattern, replacement, result, flags=re.IGNORECASE)
    return result

def create_word_document(segments, output_path, total_duration):
    """Create Word document from transcription"""
    doc = Document()
    
    title = doc.add_heading('Interrogation Transcript', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.add_run("Case: ").bold = True
    meta.add_run("Zachary Blankenship Interrogation\n")
    meta.add_run("Location: ").bold = True
    meta.add_run("Catawba County\n")
    meta.add_run("Total Duration: ").bold = True
    meta.add_run(f"{format_timestamp(total_duration)}\n")
    meta.add_run("\nParticipants:\n").bold = True
    meta.add_run("• Marcella McCombs - Lead Interrogator\n")
    meta.add_run("• Sergeant Scronce - Male Interrogator\n")
    meta.add_run("• Zachary Blankenship - Defendant\n")
    
    doc.add_paragraph()
    doc.add_paragraph("─" * 50)
    doc.add_paragraph()
    
    for segment in segments:
        timestamp = format_timestamp(segment["start"])
        text = correct_names(segment["text"])
        
        if not text.strip():
            continue
        
        p = doc.add_paragraph()
        run = p.add_run(f"[{timestamp}] ")
        run.bold = True
        run.font.size = Pt(10)
        
        text_run = p.add_run(text)
        text_run.font.size = Pt(11)
    
    doc.save(output_path)
    print(f"Saved: {output_path}")

async def main():
    audio_url = "https://customer-assets.emergentagent.com/job_legal-timeline-3/artifacts/wy0hqlmz_Audio%20Catawba.mp3"
    output_docx = "/app/backend/transcription/Catawba_Interrogation_Transcript.docx"
    json_backup = "/app/backend/transcription/transcript_segments.json"
    
    temp_dir = tempfile.mkdtemp(prefix="audio_")
    audio_path = os.path.join(temp_dir, "audio.mp3")
    
    try:
        print("Downloading...")
        subprocess.run(["curl", "-L", "-o", audio_path, audio_url], check=True, capture_output=True)
        
        print("Splitting...")
        chunks, total_duration = split_audio_ffmpeg(audio_path, temp_dir, chunk_duration=600)
        
        api_key = os.getenv("EMERGENT_LLM_KEY")
        prompt = """Police interrogation transcript. Speakers: Marcella McCombs (female lead interrogator), 
Sergeant Scronce (male interrogator), Zachary Blankenship (defendant). 
Name Rylie is spelled R-Y-L-I-E. Transcribe verbatim."""
        
        print("Transcribing...")
        all_segments = []
        
        for chunk in chunks:
            segments = await transcribe_chunk(api_key, chunk, prompt)
            all_segments.extend(segments)
            
            # Clean chunk after processing
            try:
                os.remove(chunk["path"])
            except:
                pass
        
        print(f"Total segments: {len(all_segments)}")
        
        # Save JSON backup
        with open(json_backup, 'w') as f:
            json.dump(all_segments, f, indent=2)
        print(f"JSON backup: {json_backup}")
        
        # Create Word doc
        create_word_document(all_segments, output_docx, total_duration)
        
        print("\n✓ COMPLETE!")
        
    finally:
        import shutil
        try:
            shutil.rmtree(temp_dir)
        except:
            pass

if __name__ == "__main__":
    asyncio.run(main())
